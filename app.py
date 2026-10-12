"""
PPTX Narrator V4.3

Free Piper TTS
PPTX + DOCX / PDF / TXT Speaker Notes

Features:
- Import external Speaker Notes by slide number
- Preserve slide visuals and Speaker Notes
- Complete Piper narration
- Replace existing audio
- Autoplay narration
- Advance 2 seconds after narration
- Background processing
- Batch checkpoints
- Download checkpoints
- Resume jobs while server files remain
- Recover from checkpoint/status mismatch
- Limit concurrent synthesis to one job

IMPORTANT:
Render free storage is temporary.
Download checkpoints frequently.
"""

import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
import uuid
import wave
import zipfile

from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from xml.etree import ElementTree as ET

from fastapi import (
    FastAPI,
    File,
    Form,
    HTTPException,
    UploadFile,
)

from fastapi.responses import (
    FileResponse,
    HTMLResponse,
    StreamingResponse,
)

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.package import Part


# =====================================================
# COMPATIBILITY
# =====================================================

if not hasattr(Part, "sha1"):
    Part.sha1 = property(
        lambda self: hashlib.sha1(
            self.blob
        ).hexdigest()
    )


# =====================================================
# SETTINGS
# =====================================================

app = FastAPI(title="PPTX Narrator V4.3")

ROOT = Path(__file__).resolve().parent

MODEL = (
    ROOT
    / "models"
    / "en_US-lessac-low.onnx"
)

WORKDIR = Path(
    os.environ.get(
        "NARRATOR_WORKDIR",
        str(ROOT / "narrator_jobs"),
    )
)

WORKDIR.mkdir(
    parents=True,
    exist_ok=True,
)

MAX_BYTES = 150 * 1024 * 1024
MAX_NOTES_BYTES = 25 * 1024 * 1024

PAUSE_AFTER_AUDIO_MS = 2000

P = (
    "http://schemas.openxmlformats.org/"
    "presentationml/2006/main"
)

PPTX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.presentation"
)

executor = ThreadPoolExecutor(max_workers=1)

active_jobs = set()
active_lock = threading.Lock()

# Avoid multiple large jobs waiting in memory.
MAX_ACTIVE_JOBS = 1


def ptag(name):
    return f"{{{P}}}{name}"


# =====================================================
# WEBSITE
# =====================================================

@app.get("/", response_class=HTMLResponse)
def home():

    return (
        ROOT / "index_v4.html"
    ).read_text(encoding="utf-8")


@app.get("/health")
def health():

    try:
        import fitz
        pdf_available = True
    except ImportError:
        pdf_available = False

    return {
        "status": "ok",
        "version": "4.3",
        "engine": "Piper",
        "model_exists": MODEL.is_file(),
        "external_notes": True,
        "docx_supported": True,
        "txt_supported": True,
        "pdf_supported": pdf_available,
        "api_key_required": False,
        "autoplay": True,
        "batch_processing": True,
        "storage_persistent": False,
        "active_jobs": len(active_jobs),
    }


# =====================================================
# JOB STORAGE
# =====================================================

def save_json(path, data):

    temporary = path.with_suffix(".tmp")

    temporary.write_text(
        json.dumps(
            data,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    os.replace(temporary, path)


def read_status(job_id):

    if not re.fullmatch(r"[a-f0-9]{32}", job_id):
        raise HTTPException(
            400,
            "Invalid job ID",
        )

    path = WORKDIR / job_id / "status.json"

    if not path.is_file():
        raise HTTPException(
            404,
            "Job not found. "
            "The server may have restarted. "
            "Start a new job.",
        )

    try:
        return json.loads(
            path.read_text(encoding="utf-8")
        )
    except Exception:
        raise HTTPException(
            500,
            "Job status file is damaged",
        )


def update_status(job_dir, status):

    status["updated_at"] = int(time.time())

    save_json(
        job_dir / "status.json",
        status,
    )


# =====================================================
# POWERPOINT SPEAKER NOTES
# =====================================================

def get_notes(slide):

    if not slide.has_notes_slide:
        return ""

    frame = slide.notes_slide.notes_text_frame

    if frame is None:
        return ""

    return frame.text.strip()


def set_notes(slide, text):

    frame = slide.notes_slide.notes_text_frame

    if frame is None:
        raise RuntimeError(
            "Cannot create Speaker Notes"
        )

    frame.text = text


# =====================================================
# READ TXT
# =====================================================

def read_txt(data):

    for encoding in (
        "utf-8-sig",
        "utf-16",
        "cp1252",
    ):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue

    raise RuntimeError(
        "Cannot decode TXT file"
    )


# =====================================================
# READ DOCX
# =====================================================

def read_docx(data):

    W = (
        "http://schemas.openxmlformats.org/"
        "wordprocessingml/2006/main"
    )

    def wtag(name):
        return f"{{{W}}}{name}"

    with zipfile.ZipFile(
        io.BytesIO(data)
    ) as archive:
        xml = archive.read("word/document.xml")

    root = ET.fromstring(xml)
    body = root.find(wtag("body"))

    if body is None:
        return ""

    lines = []

    for paragraph in body.iter(wtag("p")):

        pieces = []

        for node in paragraph.iter():

            if node.tag == wtag("t"):
                pieces.append(node.text or "")

            elif node.tag == wtag("tab"):
                pieces.append("\t")

            elif node.tag in (
                wtag("br"),
                wtag("cr"),
            ):
                pieces.append("\n")

        line = "".join(pieces).strip()

        if line:
            lines.append(line)

    return "\n".join(lines)


# =====================================================
# READ PDF
# =====================================================

def read_pdf(data):

    try:
        import fitz
    except ImportError:
        raise RuntimeError(
            "PDF support requires PyMuPDF"
        )

    document = fitz.open(
        stream=data,
        filetype="pdf",
    )

    try:
        return "\n".join(
            page.get_text("text", sort=True)
            for page in document
        )
    finally:
        document.close()


# =====================================================
# PARSE SLIDE HEADINGS
# =====================================================

SLIDE_HEADING = re.compile(
    r"(?im)^[ \t]*"
    r"(?:slide|幻灯片)"
    r"[ \t]*"
    r"(?:#|no\.?|number)?"
    r"[ \t]*"
    r"0*(\d{1,4})"
    r"[ \t]*"
    r"(?:[:：.\-–—][^\n]*)?"
    r"[ \t]*$"
)


def parse_slide_notes(text):

    text = text.replace(
        "\r\n", "\n"
    ).replace(
        "\r", "\n"
    )

    matches = list(
        SLIDE_HEADING.finditer(text)
    )

    if not matches:
        raise RuntimeError(
            "No slide headings found. "
            "Use Slide 001, Slide 002, etc."
        )

    notes = {}

    for index, match in enumerate(matches):

        number = int(match.group(1))

        start = match.end()

        end = (
            matches[index + 1].start()
            if index + 1 < len(matches)
            else len(text)
        )

        narration = text[start:end].strip()

        if number in notes:
            raise RuntimeError(
                f"Duplicate Slide {number}"
            )

        if not narration:
            raise RuntimeError(
                f"Slide {number} has empty narration"
            )

        notes[number] = narration

    return notes


def import_external_notes(
    prs,
    notes_data,
    filename,
):

    extension = Path(filename).suffix.lower()

    if extension == ".txt":
        text = read_txt(notes_data)

    elif extension == ".docx":
        text = read_docx(notes_data)

    elif extension == ".pdf":
        text = read_pdf(notes_data)

    else:
        raise RuntimeError(
            "Notes must be DOCX, PDF, or TXT"
        )

    mapping = parse_slide_notes(text)

    total = len(prs.slides)

    for number in mapping:

        if not 1 <= number <= total:
            raise RuntimeError(
                f"Slide {number} is outside "
                f"the PPTX range 1-{total}"
            )

    for number, narration in mapping.items():

        set_notes(
            prs.slides[number - 1],
            narration,
        )

    return len(mapping)


# =====================================================
# V4.3 STEP 1: NOTES-ONLY DOWNLOAD (NO PIPER)
# =====================================================

@app.post("/notes/insert")
async def insert_notes_only(
    file: UploadFile = File(...),
    notes_file: UploadFile = File(...),
    existing_notes: str = Form("replace"),
    require_all: bool = Form(True),
):
    """Return a PPTX with slide-numbered speaker notes, without TTS."""
    if existing_notes not in ("replace", "append", "keep"):
        raise HTTPException(400, "Invalid existing_notes mode")
    if not file.filename or Path(file.filename).suffix.lower() != ".pptx":
        raise HTTPException(400, "Please upload a PPTX")
    if not notes_file.filename or Path(notes_file.filename).suffix.lower() not in (".docx", ".pdf", ".txt"):
        raise HTTPException(400, "Please upload a DOCX, PDF, or TXT script")

    # Spool large uploads to disk rather than holding entire files in RAM.
    with tempfile.TemporaryDirectory(prefix="pptx_notes_") as tmp:
        source = Path(tmp) / "source.pptx"
        size = 0
        with source.open("wb") as out:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > MAX_BYTES:
                    raise HTTPException(413, "PPTX exceeds 150 MB")
                out.write(chunk)
        notes_data = await notes_file.read(MAX_NOTES_BYTES + 1)
        if len(notes_data) > MAX_NOTES_BYTES:
            raise HTTPException(413, "Notes exceed 25 MB")
        try:
            prs = Presentation(str(source))
            ext = Path(notes_file.filename).suffix.lower()
            script = {".txt": read_txt, ".docx": read_docx, ".pdf": read_pdf}[ext](notes_data)
            mapping = parse_slide_notes(script)
            total = len(prs.slides)
            outside = sorted(n for n in mapping if n < 1 or n > total)
            if outside:
                raise ValueError(f"Slide numbers outside PPTX (1-{total}): {outside[:12]}")
            missing = sorted(set(range(1, total + 1)) - set(mapping))
            if require_all and missing:
                raise ValueError(f"Missing narration for {len(missing)} slides: {missing[:20]}. Uncheck 'Require every slide' to allow partial scripts.")
            for n, narration in mapping.items():
                slide = prs.slides[n - 1]
                old = get_notes(slide)
                if existing_notes == "keep" and old:
                    continue
                if existing_notes == "append" and old:
                    narration = old + "\n\n" + narration
                set_notes(slide, narration)
            output = io.BytesIO()
            prs.save(output)
            output.seek(0)
            # Reopen and verify all inserted notes before returning.
            check = Presentation(output)
            for n, narration in mapping.items():
                if not get_notes(check.slides[n - 1]):
                    raise RuntimeError(f"Speaker Notes verification failed on slide {n}")
            output.seek(0)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(400, f"Notes insertion failed: {str(exc)[:450]}") from exc

    from fastapi.responses import Response
    return Response(
        content=output.getvalue(),
        media_type=PPTX_MIME,
        headers={"Content-Disposition": 'attachment; filename="PPTX_with_Speaker_Notes_V4_3.pptx"',
                 "X-Notes-Imported": str(len(mapping)),
                 "X-Total-Slides": str(total)},
    )


# =====================================================
# SPLIT NARRATION
# =====================================================

def split_narration(text, max_chars=240):

    sentences = re.split(
        r"(?<=[.!?。！？])\s+",
        text.strip(),
    )

    chunks = []
    current = ""

    for sentence in sentences:

        sentence = sentence.strip()

        if not sentence:
            continue

        if (
            len(current)
            + len(sentence)
            + 1
            <= max_chars
        ):
            current = (
                current + " " + sentence
            ).strip()
            continue

        if current:
            chunks.append(current)
            current = ""

        while len(sentence) > max_chars:

            cut = sentence.rfind(
                " ", 0, max_chars + 1
            )

            if cut <= 0:
                cut = max_chars

            part = sentence[:cut].strip()

            if part:
                chunks.append(part)

            sentence = sentence[cut:].strip()

        current = sentence

    if current:
        chunks.append(current)

    return chunks


# =====================================================
# PIPER TTS
# =====================================================

def piper_chunk(text, output_path):

    result = subprocess.run(
        [
            "piper",
            "--model",
            str(MODEL),
            "--output_file",
            str(output_path),
        ],
        input=text,
        text=True,
        capture_output=True,
        timeout=240,
    )

    if result.returncode != 0:
        raise RuntimeError(
            result.stderr[-500:]
            or "Piper generation failed"
        )

    if (
        not output_path.is_file()
        or output_path.stat().st_size < 100
    ):
        raise RuntimeError(
            "Piper produced no audio"
        )


def generate_complete_wav(text):

    if not MODEL.is_file():
        raise RuntimeError(
            "Piper model is missing"
        )

    chunks = split_narration(text)

    if not chunks:
        raise RuntimeError(
            "Empty narration"
        )

    with tempfile.TemporaryDirectory() as td:

        directory = Path(td)
        paths = []

        for index, chunk in enumerate(chunks):

            path = (
                directory
                / f"chunk_{index:04d}.wav"
            )

            piper_chunk(chunk, path)
            paths.append(path)

        output = directory / "complete.wav"

        reference = None
        total_frames = 0

        with wave.open(
            str(output), "wb"
        ) as destination:

            for path in paths:

                with wave.open(
                    str(path), "rb"
                ) as source:

                    params = (
                        source.getnchannels(),
                        source.getsampwidth(),
                        source.getframerate(),
                        source.getcomptype(),
                    )

                    if reference is None:

                        reference = params

                        destination.setnchannels(
                            params[0]
                        )

                        destination.setsampwidth(
                            params[1]
                        )

                        destination.setframerate(
                            params[2]
                        )

                    elif params != reference:
                        raise RuntimeError(
                            "WAV formats differ"
                        )

                    frames = source.readframes(
                        source.getnframes()
                    )

                    if not frames:
                        raise RuntimeError(
                            "Empty WAV chunk"
                        )

                    destination.writeframes(frames)

                    total_frames += source.getnframes()

        duration_ms = round(
            1000 * total_frames / reference[2]
        )

        return (
            output.read_bytes(),
            duration_ms,
            len(chunks),
        )


# =====================================================
# REMOVE OLD AUDIO
# =====================================================

def remove_old_audio(slide):

    for shape in list(slide.shapes):

        if (
            shape.shape_type
            == MSO_SHAPE_TYPE.MEDIA
        ):
            slide.shapes._spTree.remove(
                shape._element
            )

    # This also removes existing animations.
    for timing in list(
        slide._element.findall(ptag("timing"))
    ):
        slide._element.remove(timing)


# =====================================================
# INSERT AUDIO
# =====================================================

def insert_wav(slide, audio_bytes):

    with tempfile.NamedTemporaryFile(
        suffix=".wav",
        delete=False,
    ) as temporary:

        temporary.write(audio_bytes)
        path = Path(temporary.name)

    try:

        shape = slide.shapes.add_movie(
            str(path),
            left=0,
            top=0,
            width=320040,
            height=320040,
            mime_type="audio/wav",
        )

        return shape.shape_id

    finally:
        path.unlink(missing_ok=True)


# =====================================================
# AUTOPLAY
# =====================================================

def configure_autoplay(
    slide,
    shape_id,
    duration_ms,
    advance=True,
):

    for old in list(
        slide._element.findall(ptag("timing"))
    ):
        slide._element.remove(old)

    timing = etree.Element(ptag("timing"))

    tn_list = etree.SubElement(
        timing, ptag("tnLst")
    )

    par = etree.SubElement(
        tn_list, ptag("par")
    )

    root = etree.SubElement(
        par,
        ptag("cTn"),
        id="1",
        dur="indefinite",
        restart="never",
        nodeType="tmRoot",
    )

    children = etree.SubElement(
        root,
        ptag("childTnLst"),
    )

    audio = etree.SubElement(
        children,
        ptag("audio"),
    )

    media = etree.SubElement(
        audio,
        ptag("cMediaNode"),
        vol="80000",
    )

    media_timing = etree.SubElement(
        media,
        ptag("cTn"),
        id="2",
        fill="hold",
        display="0",
    )

    conditions = etree.SubElement(
        media_timing,
        ptag("stCondLst"),
    )

    etree.SubElement(
        conditions,
        ptag("cond"),
        delay="0",
    )

    target = etree.SubElement(
        media,
        ptag("tgtEl"),
    )

    etree.SubElement(
        target,
        ptag("spTgt"),
        spid=str(shape_id),
    )

    extension = slide._element.find(
        ptag("extLst")
    )

    if extension is not None:
        slide._element.insert(
            slide._element.index(extension),
            timing,
        )
    else:
        slide._element.append(timing)

    transition = slide._element.find(
        ptag("transition")
    )

    if transition is None:

        transition = etree.Element(
            ptag("transition")
        )

        slide._element.insert(
            slide._element.index(timing),
            transition,
        )

    if advance:

        transition.set("advClick", "0")

        transition.set(
            "advTm",
            str(
                duration_ms
                + PAUSE_AFTER_AUDIO_MS
            ),
        )

    else:

        transition.set("advClick", "1")

        transition.attrib.pop(
            "advTm", None
        )


# =====================================================
# VALIDATE PPTX
# =====================================================

def validate_pptx(path, expected_slides):

    with zipfile.ZipFile(path) as archive:

        if archive.testzip() is not None:
            raise RuntimeError(
                "PPTX ZIP validation failed"
            )

    check = Presentation(str(path))

    if len(check.slides) != expected_slides:
        raise RuntimeError(
            "Slide count changed"
        )


# =====================================================
# CHECKPOINT STORAGE
# =====================================================

def checkpoint_path(job_dir):
    return job_dir / "checkpoint.pptx"


def checkpoint_meta_path(job_dir):
    return job_dir / "checkpoint.json"


def save_checkpoint(
    prs,
    job_dir,
    total,
    completed,
    skipped,
):

    temporary = (
        job_dir / "checkpoint.tmp.pptx"
    )

    destination = checkpoint_path(job_dir)

    prs.save(str(temporary))

    validate_pptx(temporary, total)

    os.replace(temporary, destination)

    save_json(
        checkpoint_meta_path(job_dir),
        {
            "completed": completed,
            "skipped_empty_notes": skipped,
            "total": total,
            "saved_at": int(time.time()),
        },
    )


def restore_checkpoint_status(job_dir, status):

    checkpoint = checkpoint_path(job_dir)
    metadata = checkpoint_meta_path(job_dir)

    if not checkpoint.is_file():
        return status

    if not metadata.is_file():
        return status

    try:

        saved = json.loads(
            metadata.read_text(
                encoding="utf-8"
            )
        )

        completed = int(
            saved["completed"]
        )

        if (
            0 <= completed <= status["target"]
        ):
            status["completed"] = completed

            status["skipped_empty_notes"] = int(
                saved.get(
                    "skipped_empty_notes", 0
                )
            )

            status["checkpoint_available"] = True

    except Exception:
        pass

    return status


# =====================================================
# BACKGROUND PROCESSING
# =====================================================

def process_job(job_id):

    job_dir = WORKDIR / job_id

    try:
        status = read_status(job_id)
    except HTTPException:
        with active_lock:
            active_jobs.discard(job_id)
        return

    try:

        status = restore_checkpoint_status(
            job_dir,
            status,
        )

        status["state"] = "running"
        status["error"] = None

        update_status(job_dir, status)

        total = status["total"]
        target = status["target"]
        batch_size = status["batch_size"]

        checkpoint = checkpoint_path(job_dir)

        source = (
            checkpoint
            if checkpoint.is_file()
            else job_dir / "input.pptx"
        )

        prs = Presentation(str(source))

        completed = int(status["completed"])

        for i in range(completed, target):

            status["current_slide"] = i + 1

            update_status(job_dir, status)

            slide = prs.slides[i]
            note = get_notes(slide)

            if note:

                (
                    audio_bytes,
                    duration_ms,
                    chunk_count,
                ) = generate_complete_wav(note)

                remove_old_audio(slide)

                shape_id = insert_wav(
                    slide,
                    audio_bytes,
                )

                configure_autoplay(
                    slide,
                    shape_id,
                    duration_ms,
                    advance=(i < total - 1),
                )

                status["last_slide_seconds"] = round(
                    duration_ms / 1000, 1
                )

                status["last_slide_chunks"] = (
                    chunk_count
                )

                del audio_bytes

            else:
                status["skipped_empty_notes"] += 1

            if (
                (i + 1) % batch_size == 0
                or i + 1 == target
            ):

                save_checkpoint(
                    prs,
                    job_dir,
                    total,
                    i + 1,
                    status["skipped_empty_notes"],
                )

                status["completed"] = i + 1

                status["checkpoint_available"] = True

            update_status(job_dir, status)

        status["state"] = "completed"
        status["current_slide"] = target

        update_status(job_dir, status)

    except Exception as exc:

        status["state"] = "failed"

        status["error"] = (
            f"Slide "
            f"{status.get('current_slide', 0)}: "
            f"{str(exc)[:400]}"
        )

        try:
            update_status(job_dir, status)
        except Exception:
            pass

    finally:

        with active_lock:
            active_jobs.discard(job_id)


def launch_job(job_id):

    with active_lock:

        if job_id in active_jobs:
            return False

        if len(active_jobs) >= MAX_ACTIVE_JOBS:
            raise HTTPException(
                429,
                "Another narration job is running. "
                "Please wait.",
            )

        active_jobs.add(job_id)

    try:
        executor.submit(
            process_job,
            job_id,
        )
    except Exception:
        with active_lock:
            active_jobs.discard(job_id)
        raise

    return True


# =====================================================
# CREATE JOB
# =====================================================

@app.post("/jobs")
async def create_job(
    file: UploadFile = File(...),
    notes_file: UploadFile = File(None),
    limit: int = Form(0),
    batch_size: int = Form(5),
):

    if (
        not file.filename
        or not file.filename.lower().endswith(
            ".pptx"
        )
    ):
        raise HTTPException(
            400,
            "Upload a PPTX file",
        )

    if limit < 0:
        raise HTTPException(
            400,
            "Invalid slide limit",
        )

    if not 1 <= batch_size <= 20:
        raise HTTPException(
            400,
            "Batch size must be 1-20",
        )

    if not MODEL.is_file():
        raise HTTPException(
            500,
            "Piper model is missing",
        )

    # Reject a new job before processing upload
    # if another job is already running.
    with active_lock:
        if len(active_jobs) >= MAX_ACTIVE_JOBS:
            raise HTTPException(
                429,
                "Another narration job is running. "
                "Please wait.",
            )

    job_id = uuid.uuid4().hex
    job_dir = WORKDIR / job_id

    job_dir.mkdir(
        parents=True,
        exist_ok=False,
    )

    source = job_dir / "input.pptx"

    try:

        size = 0

        with source.open("wb") as output:

            while True:

                chunk = await file.read(
                    1024 * 1024
                )

                if not chunk:
                    break

                size += len(chunk)

                if size > MAX_BYTES:
                    raise HTTPException(
                        413,
                        "PPTX exceeds 150 MB",
                    )

                output.write(chunk)

        try:
            prs = Presentation(str(source))
        except Exception:
            raise HTTPException(
                400,
                "Invalid PPTX",
            )

        total = len(prs.slides)

        if total == 0:
            raise HTTPException(
                400,
                "PPTX contains no slides",
            )

        target = (
            min(limit, total)
            if limit
            else total
        )

        imported_count = 0

        if (
            notes_file is not None
            and notes_file.filename
        ):

            notes_data = await notes_file.read(
                MAX_NOTES_BYTES + 1
            )

            if len(notes_data) > MAX_NOTES_BYTES:
                raise HTTPException(
                    413,
                    "Notes file exceeds 25 MB",
                )

            try:

                imported_count = import_external_notes(
                    prs,
                    notes_data,
                    notes_file.filename,
                )

            except Exception as exc:
                raise HTTPException(
                    400,
                    "Notes import failed: "
                    + str(exc)[:400],
                )

            prs.save(str(source))

        notes_found = sum(
            bool(get_notes(prs.slides[i]))
            for i in range(target)
        )

        if notes_found == 0:
            raise HTTPException(
                400,
                "No Speaker Notes found. "
                "Upload a matching DOCX, PDF, or TXT.",
            )

        status = {
            "job_id": job_id,
            "state": "queued",
            "total": total,
            "target": target,
            "completed": 0,
            "current_slide": 0,
            "batch_size": batch_size,
            "imported_notes": imported_count,
            "notes_in_target": notes_found,
            "skipped_empty_notes": 0,
            "checkpoint_available": False,
            "error": None,
            "created_at": int(time.time()),
        }

        update_status(job_dir, status)

        launch_job(job_id)

        return status

    except Exception:

        shutil.rmtree(
            job_dir,
            ignore_errors=True,
        )

        raise


# =====================================================
# JOB STATUS
# =====================================================

@app.get("/jobs/{job_id}")
def job_status(job_id: str):

    status = read_status(job_id)

    job_dir = WORKDIR / job_id

    status = restore_checkpoint_status(
        job_dir,
        status,
    )

    with active_lock:
        is_active = job_id in active_jobs

    if (
        status["state"] in ("running", "queued")
        and not is_active
    ):
        status["state"] = "interrupted"

    status["download_url"] = (
        f"/jobs/{job_id}/download"
        if checkpoint_path(job_dir).is_file()
        else None
    )

    return status


# =====================================================
# RESUME JOB
# =====================================================

@app.post("/jobs/{job_id}/resume")
def resume_job(job_id: str):

    status = read_status(job_id)

    if status["state"] == "completed":
        return status

    with active_lock:
        if job_id in active_jobs:
            raise HTTPException(
                409,
                "Job is already running",
            )

    status["state"] = "queued"
    status["error"] = None

    update_status(
        WORKDIR / job_id,
        status,
    )

    launch_job(job_id)

    return status


# =====================================================
# DOWNLOAD CHECKPOINT
# =====================================================

@app.get("/jobs/{job_id}/download")
def download_job(job_id: str):

    status = read_status(job_id)

    job_dir = WORKDIR / job_id

    status = restore_checkpoint_status(
        job_dir,
        status,
    )

    path = checkpoint_path(job_dir)

    if not path.is_file():
        raise HTTPException(
            404,
            "No checkpoint available",
        )

    return FileResponse(
        str(path),
        media_type=PPTX_MIME,
        filename=(
            f"narrated_v4_2_"
            f"{status['completed']}_"
            f"of_{status['total']}.pptx"
        ),
    )


# =====================================================
# V3 COMPATIBILITY
# =====================================================

@app.post("/process")
async def legacy_process(
    file: UploadFile = File(...),
    mode: str = Form("tts"),
    voice: str = Form("alloy"),
    api_key: str = Form(""),
    limit: int = Form(0),
):

    raw = await file.read(MAX_BYTES + 1)

    if len(raw) > MAX_BYTES:
        raise HTTPException(
            413,
            "PPTX exceeds 150 MB",
        )

    try:
        prs = Presentation(io.BytesIO(raw))
    except Exception:
        raise HTTPException(
            400,
            "Invalid PPTX",
        )

    total = len(prs.slides)

    if limit < 0:
        raise HTTPException(
            400,
            "Invalid slide limit",
        )

    count = (
        min(limit, total)
        if limit
        else total
    )

    notes = [
        get_notes(prs.slides[i])
        for i in range(count)
    ]

    if mode == "notes":

        text = "\n\n".join(
            f"=== Slide {i + 1} ===\n{note}"
            for i, note in enumerate(notes)
        )

        return StreamingResponse(
            io.BytesIO(
                text.encode("utf-8")
            ),
            media_type="text/plain; charset=utf-8",
            headers={
                "Content-Disposition":
                'attachment; filename="slide_notes.txt"'
            },
        )

    if mode == "existing":

        return StreamingResponse(
            io.BytesIO(raw),
            media_type=PPTX_MIME,
            headers={
                "Content-Disposition":
                'attachment; filename="original_audio.pptx"'
            },
        )

    if mode != "tts":
        raise HTTPException(
            400,
            "Invalid mode",
        )

    if not any(notes):
        raise HTTPException(
            400,
            "No Speaker Notes found",
        )

    for i, note in enumerate(notes):

        if not note:
            continue

        try:

            audio, duration, _ = (
                generate_complete_wav(note)
            )

            slide = prs.slides[i]

            remove_old_audio(slide)

            shape_id = insert_wav(
                slide, audio
            )

            configure_autoplay(
                slide,
                shape_id,
                duration,
                advance=(i < total - 1),
            )

        except Exception as exc:
            raise HTTPException(
                500,
                f"Slide {i + 1}: {str(exc)[:300]}",
            )

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=PPTX_MIME,
        headers={
            "Content-Disposition":
            'attachment; filename="narrated_piper_v4_2.pptx"'
        },
    )
